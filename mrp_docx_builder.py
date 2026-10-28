"""
Builds the 21L MRP label sheet (see mrp_label_builder.py) as a Word
.docx -- same geometry and layout as the PDF, but with real editable text
and tables, so it can be opened and printed straight from Word.

Word lays text out itself and never shrinks it to fit, so the font size of
each label's details table is chosen up front with the PDF builder's fit
logic (Times New Roman and reportlab's Times share metrics), with a little safety margin.
"""

import io

import docx
from docx.enum.table import WD_ALIGN_VERTICAL, WD_ROW_HEIGHT_RULE, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Mm, Pt

from label_builder import (
    PAGE_WIDTH_MM, PAGE_HEIGHT_MM,
    make_barcode_stream, _set_no_spacing, _FastPicInserter,
)
from mrp_label_builder import (
    TOP_MARGIN_MM, SIDE_MARGIN_MM, LABEL_WIDTH_MM, LABEL_HEIGHT_MM,
    HORIZONTAL_PITCH_MM, VERTICAL_PITCH_MM, COLS_ACROSS, ROWS_DOWN, LABELS_PER_PAGE,
    INNER_PAD_MM, BOTTOM_PAD_MM, BARCODE_GAP_MM, TABLE_TOP_MM,
    BARCODE_WIDTH_MM, BARCODE_HEIGHT_MM,
    LEFT_COL_FRACTION, FONT_SIZES_PT, LINE_FACTOR, CELL_PAD_X_MM, CELL_PAD_Y_MM,
    IMPORTER_ADDRESS, IMPORTER_CONTACT, QUANTITY_TEXT, _build_table,
)
from reportlab.lib.units import mm as rl_mm

FONT = "Times New Roman"
GAP_WIDTH_MM = HORIZONTAL_PITCH_MM - LABEL_WIDTH_MM          # 2.5 mm between columns
TRAILING_PARA_PT = 1                                         # the paragraph after each nested table
# Height the details table may use in Word: label height minus what sits
# above it (the bars) and the bottom margin. Word draws the table ~0.3 mm
# taller than reportlab measures it (checked via Word -> PDF).
WORD_TABLE_EXTRA_MM = 0.3
WORD_TABLE_BUDGET_MM = LABEL_HEIGHT_MM - TABLE_TOP_MM - BOTTOM_PAD_MM - WORD_TABLE_EXTRA_MM

CAPTIONS = [
    "Name and address of Importer & Marketed by & Contact if Consumer Complaints",
    "Product Name / Model ID / Colour / Quantity / Month of Import / MRP (Inclusive of All Taxes)",
    "Manufacturers Address & Country of Origin",
    "Year of Mfg",
]


def _mm_to_twips(v):
    return str(int(round(v / 25.4 * 1440)))


# Word only honours tcPr children in schema order (an out-of-order vAlign or
# tcMar is silently ignored), so every cell property goes through _put_tcpr.
TCPR_ORDER = ["cnfStyle", "tcW", "gridSpan", "hMerge", "vMerge", "tcBorders", "shd",
              "noWrap", "tcMar", "textDirection", "tcFitText", "vAlign", "hideMark"]


def _put_tcpr(cell, el):
    """Set one tcPr child, replacing any existing one, at its schema position."""
    tcPr = cell._tc.get_or_add_tcPr()
    tag = el.tag
    for old in tcPr.findall(tag):
        tcPr.remove(old)
    rank = TCPR_ORDER.index(tag.split("}")[1])
    for child in tcPr:
        name = child.tag.split("}")[1]
        if name in TCPR_ORDER and TCPR_ORDER.index(name) > rank:
            child.addprevious(el)
            return
    tcPr.append(el)


def _set_cell_margins_mm(cell, top=0.0, bottom=0.0, left=0.0, right=0.0):
    mar = OxmlElement("w:tcMar")
    for tag, val in (("top", top), ("bottom", bottom), ("start", left), ("end", right)):
        node = OxmlElement(f"w:{tag}")
        node.set(qn("w:w"), _mm_to_twips(val))
        node.set(qn("w:type"), "dxa")
        mar.append(node)
    _put_tcpr(cell, mar)


def _set_table_borders(table, val="single", size_eighths=2):
    tbl_pr = table._tbl.tblPr
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        b = OxmlElement(f"w:{edge}")
        b.set(qn("w:val"), val)
        if val != "nil":
            b.set(qn("w:sz"), str(size_eighths))
            b.set(qn("w:color"), "000000")
        borders.append(b)
    tbl_pr.append(borders)


def _zero_table_offsets(table):
    """Word indents a table and pads its cells by default, which pushed the
    whole grid ~1.5 mm left of the 7 mm side margin -- zero both."""
    tbl_pr = table._tbl.tblPr
    ind = OxmlElement("w:tblInd")
    ind.set(qn("w:w"), "0")
    ind.set(qn("w:type"), "dxa")
    # Word reads tblPr children in schema order: tblInd follows tblW and
    # tblCellMar precedes tblLook, else they are ignored
    tbl_pr.find(qn("w:tblW")).addnext(ind)
    mar = OxmlElement("w:tblCellMar")
    for tag in ("top", "left", "start", "bottom", "right", "end"):
        node = OxmlElement(f"w:{tag}")
        node.set(qn("w:w"), "0")
        node.set(qn("w:type"), "dxa")
        mar.append(node)
    tbl_pr.find(qn("w:tblLook")).addprevious(mar)


def _set_cell_border(cell, color="9E9E9E"):
    borders = OxmlElement("w:tcBorders")
    for edge in ("top", "left", "bottom", "right"):
        b = OxmlElement(f"w:{edge}")
        b.set(qn("w:val"), "single")
        b.set(qn("w:sz"), "2")
        b.set(qn("w:color"), color)
        borders.append(b)
    _put_tcpr(cell, borders)


def _shade(cell, fill="D9D9D9"):
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill)
    _put_tcpr(cell, shd)


def _fixed_layout(table):
    layout = OxmlElement("w:tblLayout")
    layout.set(qn("w:type"), "fixed")
    table._tbl.tblPr.append(layout)


def _exact_lines(paragraph, font_pt):
    _set_no_spacing(paragraph)
    pf = paragraph.paragraph_format
    pf.line_spacing_rule = WD_LINE_SPACING.EXACTLY
    pf.line_spacing = Pt(round(font_pt * LINE_FACTOR * 4) / 4)


def _add_text(paragraph, text, font_pt, bold=False):
    lines = str(text).split("\n")
    for i, line in enumerate(lines):
        run = paragraph.add_run(line)
        run.font.name = FONT
        run.font.size = Pt(font_pt)
        run.font.bold = bold
        if i < len(lines) - 1:
            run.add_break()


def _details_font_pt(info, width_mm):
    """Largest font whose table fits the Word budget (same metrics as the PDF builder)."""
    for font_pt in FONT_SIZES_PT:
        _, h = _build_table(info, font_pt, width_mm).wrap(width_mm * rl_mm, WORD_TABLE_BUDGET_MM * rl_mm)
        if h <= WORD_TABLE_BUDGET_MM * rl_mm:
            return font_pt
    return FONT_SIZES_PT[-1]


def _fill_label(cell, info, barcode_stream, pic_inserter):
    inner_w = LABEL_WIDTH_MM - 2 * INNER_PAD_MM
    _set_cell_margins_mm(cell, left=INNER_PAD_MM, right=INNER_PAD_MM)
    # Word ignores vertical centring in a cell holding a nested table, so the
    # block (bars + gap + table) is centred by hand: the space above the bars
    # is computed from the table's measured height, same as the PDF preview.
    cell.vertical_alignment = WD_ALIGN_VERTICAL.TOP
    font_pt = _details_font_pt(info, inner_w)
    _, table_h = _build_table(info, font_pt, inner_w).wrap(inner_w * rl_mm, LABEL_HEIGHT_MM * rl_mm)
    block_mm = (BARCODE_HEIGHT_MM + BARCODE_GAP_MM + table_h / rl_mm
                + WORD_TABLE_EXTRA_MM + TRAILING_PARA_PT * 25.4 / 72)
    top_mm = max(0.0, (LABEL_HEIGHT_MM - block_mm) / 2)

    # bars (no barcode number)
    p = cell.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    # single (not exact) spacing: an exact line shorter than the picture
    # would crop the bars
    _set_no_spacing(p)
    p.paragraph_format.space_before = Mm(top_mm)
    p.paragraph_format.space_after = Mm(BARCODE_GAP_MM)
    barcode_stream.seek(0)
    pic_inserter.add_picture(p, barcode_stream, Mm(BARCODE_WIDTH_MM), Mm(BARCODE_HEIGHT_MM))

    # details table
    values = [
        f"{info['importer']}, {IMPORTER_ADDRESS}\n{IMPORTER_CONTACT}",
        "\n".join([info["product_type"], info["model_name"], QUANTITY_TEXT,
                   info["month_of_import"], f"Rs.{info['mrp']}/-"]),
        info["manufacturer"],
        info["year_of_mfg"],
    ]
    left_w = inner_w * LEFT_COL_FRACTION
    widths = [Mm(left_w), Mm(inner_w - left_w)]
    t = cell.add_table(rows=len(CAPTIONS), cols=2)
    t.autofit = False
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    _fixed_layout(t)
    _set_table_borders(t)
    for col, w in zip(t.columns, widths):
        col.width = w
    for r, (cap, val) in enumerate(zip(CAPTIONS, values)):
        for c, text in enumerate((cap, val)):
            tc = t.rows[r].cells[c]
            tc.width = widths[c]
            tc.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            _set_cell_margins_mm(tc, top=CELL_PAD_Y_MM, bottom=CELL_PAD_Y_MM,
                                 left=CELL_PAD_X_MM, right=CELL_PAD_X_MM)
            tp = tc.paragraphs[0]
            tp.alignment = WD_ALIGN_PARAGRAPH.CENTER
            _exact_lines(tp, font_pt)
            _add_text(tp, text, font_pt)
    _shade(t.rows[0].cells[0])

    # Word requires a paragraph after a nested table -- keep it ~invisible
    p = cell.add_paragraph()
    _set_no_spacing(p)
    p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.EXACTLY
    p.paragraph_format.line_spacing = Pt(TRAILING_PARA_PT)


def _new_sheet_table(document, vertical_pitch_mm, show_outlines):
    """One page: 7 rows x (label, gap, label, gap, label), each row one pitch high."""
    table = document.add_table(rows=ROWS_DOWN, cols=COLS_ACROSS * 2 - 1)
    table.autofit = False
    _fixed_layout(table)
    _set_table_borders(table, val="nil")
    _zero_table_offsets(table)
    widths = [Mm(LABEL_WIDTH_MM) if i % 2 == 0 else Mm(GAP_WIDTH_MM)
              for i in range(COLS_ACROSS * 2 - 1)]
    for i, col in enumerate(table.columns):
        col.width = widths[i]
    for row in table.rows:
        row.height_rule = WD_ROW_HEIGHT_RULE.EXACTLY
        # Measured via Word -> PDF: unlike the 84L sheet, these rows need
        # no WORD_ROW_BOUNDARY_PT correction -- the pitch comes out exact.
        row.height = Mm(vertical_pitch_mm)
        for i, cell in enumerate(row.cells):
            cell.width = widths[i]
            _set_cell_margins_mm(cell)
            _set_no_spacing(cell.paragraphs[0])
            if show_outlines and i % 2 == 0:
                _set_cell_border(cell)
    return table


def build_mrp_label_docx(entries, start_position=1, top_margin_mm=TOP_MARGIN_MM,
                         left_margin_mm=SIDE_MARGIN_MM,
                         vertical_pitch_mm=VERTICAL_PITCH_MM,
                         show_outlines=False) -> io.BytesIO:
    """Same arguments as mrp_label_builder.build_mrp_label_pdf; returns a .docx (BytesIO)."""
    document = docx.Document()
    # python-docx's template is Word 2010 mode (14), where Word pulls each
    # table left by its first cell's padding -- the 1.5 mm label pad put the
    # whole grid 1.5 mm left of the die-cuts. Mode 15 keeps it on the margin.
    for cs in document.settings.element.iter(qn("w:compatSetting")):
        if cs.get(qn("w:name")) == "compatibilityMode":
            cs.set(qn("w:val"), "15")
    section = document.sections[0]
    section.page_width = Mm(PAGE_WIDTH_MM)
    section.page_height = Mm(PAGE_HEIGHT_MM)
    section.top_margin = Mm(top_margin_mm)
    section.bottom_margin = Mm(0)
    section.left_margin = Mm(left_margin_mm)
    section.right_margin = Mm(max(0.0, PAGE_WIDTH_MM - left_margin_mm
                                  - (COLS_ACROSS - 1) * HORIZONTAL_PITCH_MM - LABEL_WIDTH_MM))
    section.header_distance = Mm(0)
    section.footer_distance = Mm(0)

    img_cache = {}
    for info in entries:
        if info["barcode"] not in img_cache:
            img_cache[info["barcode"]] = make_barcode_stream(info["barcode"])
    pic_inserter = _FastPicInserter(document)

    # first slot used on page 1; later pages start at slot 0
    slot = max(1, int(start_position)) - 1
    idx, total = 0, len(entries)
    first_page = True
    while first_page or idx < total:
        if not first_page:
            document.add_page_break()
        table = _new_sheet_table(document, vertical_pitch_mm, show_outlines)
        while slot < LABELS_PER_PAGE and idx < total:
            row, col = divmod(slot, COLS_ACROSS)
            _fill_label(table.rows[row].cells[col * 2], entries[idx],
                        img_cache[entries[idx]["barcode"]], pic_inserter)
            slot += 1
            idx += 1
        slot = 0
        first_page = False

    # the paragraph python-docx leaves after the last table must not spill a blank page
    last = document.paragraphs[-1] if document.paragraphs else None
    if last is not None:
        _set_no_spacing(last)
        last.paragraph_format.line_spacing_rule = WD_LINE_SPACING.EXACTLY
        last.paragraph_format.line_spacing = Pt(1)

    out = io.BytesIO()
    document.save(out)
    out.seek(0)
    return out
